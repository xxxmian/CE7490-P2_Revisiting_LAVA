#!/usr/bin/env python3
"""Build the course deck from recorded results (requires: pip install python-pptx).
Uses the previous deck as the template (masters, theme, two illustrations) and regenerates
every slide in its visual style. All numbers are read from results/ or recomputed."""
import argparse,csv,io,json,math,statistics,sys
from pathlib import Path
from lxml import etree
root=Path(__file__).resolve().parents[1];sys.path.insert(0,str(root))
from lava.model import BayesPredictor,Predictor,Request,mixture_remaining,trace
from lava.stats import paired_diffs,t_test
from pptx import Presentation
from pptx.util import Inches,Pt,Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR,MSO_AUTO_SIZE
from pptx.chart.data import CategoryChartData,XyChartData
from pptx.enum.chart import XL_CHART_TYPE,XL_LEGEND_POSITION,XL_MARKER_STYLE,XL_TICK_LABEL_POSITION,XL_LABEL_POSITION
from pptx.oxml.ns import qn

# Palette of the original deck. Orange text uses a darker step for contrast on the beige background.
NAVY,ORANGE,ORANGE_TXT,GREEN,RED,INK,MUTED,CAPTION,BG,EDGE,SLATE='2F5D8A','D9822B','A85A14','3E7C59','B64A4A','233142','6B7785','5A6573','F7F4EE','E7EBEF','6B7785'
POS,NEG,FONT='BFE3C9','F2C4C1','Aptos'
TXT={ORANGE:ORANGE_TXT}
REVISED='29 October 2026';FOOT=f'Independent reconstruction   {REVISED}';NB=' '

# ---------- data (same sources as build_report.py) ----------
RES=root/'results';J=lambda p:json.loads((RES/p).read_text())
C=json.loads((root/'configs/report_scale.json').read_text());PROTO=json.loads((root/'configs/improvement_protocol.json').read_text())
SUITES=['test','legacy','tau','robust','robust_control','robust_load']
IMP={s:J(f'improvement/{s}/summary.json') for s in SUITES};MAN={s:J(f'improvement/{s}/manifest.json') for s in SUITES}
HELD,SWEEP=J('heldout/summary.json'),J('sweep/summary.json')
GB={(r['setting'],r['confidence']):r for r in J('improvement/gate_by_bucket.json')['rows']}
HIST={(r['setting'],r['policy']):r for r in csv.DictReader((root/'reference/report_table1.csv').open())}
ACC=['0.25','0.5','0.75','0.85','1.0'];REG=ACC[:4]
def pick(rows,policy,baseline,setting,metric='empty_pct',**kw):
    return next(r for r in rows if r['policy']==policy and r['baseline']==baseline and r['setting']==setting and r.get('metric','empty_pct')==metric and all(r[k]==v for k,v in kw.items()))
def num(x,d=3):return f'{0:.{d}f}' if round(x,d)==0 else f'{x:+.{d}f}'.replace('-','−')
def ci(r,d=3):return f"{num(r['mean'],d)} [{num(r['low'],d)}, {num(r['high'],d)}]"
def pf(p):return '<0.001' if p<1e-3 else (f'{p:.3f}' if p<.01 else f'{p:.2f}')
def lab(s):return f'{float(s):.0%}' if s[0].isdigit() else ('Default' if s=='synthetic' else s.capitalize())
TEST_RAW=[dict(r,seed=int(r['seed']),empty_pct=float(r['empty_pct']),empty_pct_arrivals=float(r['empty_pct_arrivals']),accepted=int(r['accepted']),failures=int(r['failures']))
          for r in csv.DictReader((RES/'improvement/test/raw.csv').open())]
IDX={(r['seed'],r['setting'],r['policy']):r for r in TEST_RAW};SEEDS=sorted({r['seed'] for r in TEST_RAW})
HELD_RAW=list(csv.DictReader((RES/'heldout/raw.csv').open()))
def res(m,lo,hi,p):return dict(mean=m,low=lo,high=hi,p=p)
def posthoc(p,b,s):return res(*t_test(paired_diffs(TEST_RAW,p,b,s)))
def cross(p,s,b,bs):return res(*t_test([IDX[(k,s,p)]['empty_pct']-IDX[(k,bs,b)]['empty_pct'] for k in SEEDS]))
def gate_rate(s,p):return 100*statistics.mean(int(r['high_gate'])/r['accepted'] for r in TEST_RAW if r['setting']==s and r['policy']==p)
def held_mean(p,f):return statistics.mean(float(r[f]) for r in HELD_RAW if r['policy']==p and r['setting']=='synthetic')
def tail_recovery():
    """Share of Host-CA's lost empty host-hours (arrival window) that the drain tail recovers."""
    lost_a=lost_f=0.
    for k in SEEDS:
        recs=trace(C,k);t0=recs[0].request.arrival;ta=max(r.request.arrival for r in recs);bf,hc=IDX[(k,'synthetic','best_fit')],IDX[(k,'synthetic','host_ca')]
        lost_a+=(bf['empty_pct_arrivals']-hc['empty_pct_arrivals'])/100*C['hosts']*(ta-t0);lost_f+=(bf['empty_pct']-hc['empty_pct'])/100*C['hosts']*float(bf['horizon'])
    return 1-lost_f/lost_a
FAMILY={(r['setting'],r['baseline']):r for r in IMP['test'] if r['primary']};ok=lambda r:r['mean']>0 and r['p_holm']<.05
HOLDS=sum(ok(r) for r in FAMILY.values());HOLDS_LAVA=sum(ok(FAMILY[(a,'lava')]) for a in REG);HOLDS_BF=sum(ok(FAMILY[(a,'best_fit')]) for a in REG)
IMP_RUNS=sum(MAN[s]['runs'] for s in SUITES);IMP_MIN=sum(MAN[s]['elapsed_s'] for s in SUITES)/60;N_TEST=MAN['test']['n_seeds']
T=lambda p,b,s,*a,**kw:pick(IMP['test'],p,b,s,*a,**kw)
def profile(name):
    p,a,b=C['profiles'][name];mean,w=mixture_remaining((p,a,b),0);edges=[0]+C['class_bounds']+[math.inf]
    Sv=lambda t:0. if t==math.inf else w*math.exp(-t/a)+(1-w)*math.exp(-t/b)
    pr=[Sv(lo)-Sv(hi) for lo,hi in zip(edges,edges[1:])];k=sum(1 for x in C['class_bounds'] if mean>=x);return pr,mean,k
def predict_pair(profile_name,setting,value,age=0):
    q=BayesPredictor(C,0,setting,[]);q.values={0:value};r=Predictor(C,0,setting,[]);r.values={0:value}
    req=Request(0,0,.125,.125,profile_name);return q.predict(req,age),r.predict(req,age)

# ---------- drawing helpers ----------
def rgb(h):return RGBColor.from_string(h)
def style_run(r,size,color,bold,italic=False):
    f=r.font;f.size=Pt(size);f.bold=bold;f.italic=italic;f.color.rgb=rgb(color);f.name=FONT;rPr=r._r.get_or_add_rPr()
    for tag in ('a:ea','a:cs'):
        el=rPr.find(qn(tag))
        if el is None:el=etree.SubElement(rPr,qn(tag))
        el.set('typeface',FONT)
def text(s,x,y,w,h,paras,size,color=INK,bold=False,align='l',anchor='t',space=0):
    tb=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=tb.text_frame;tf.word_wrap=True;tf.auto_size=MSO_AUTO_SIZE.NONE
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=Emu(36576);tf.vertical_anchor={'t':MSO_ANCHOR.TOP,'m':MSO_ANCHOR.MIDDLE}[anchor]
    for i,para in enumerate([paras] if isinstance(paras,str) else paras):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.alignment={'l':PP_ALIGN.LEFT,'c':PP_ALIGN.CENTER,'r':PP_ALIGN.RIGHT}[align]
        if space and i:p.space_before=Pt(space)
        for t,o in (para if isinstance(para,list) else [(para,{})]):
            r=p.add_run();r.text=t;style_run(r,o.get('size',size),o.get('color',color),o.get('bold',bold),o.get('italic',False))
    return tb
def box(s,x,y,w,h,fill='FFFFFF',line=EDGE,shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    b=s.shapes.add_shape(shape,Inches(x),Inches(y),Inches(w),Inches(h));b.fill.solid();b.fill.fore_color.rgb=rgb(fill)
    if shape==MSO_SHAPE.ROUNDED_RECTANGLE:b.adjustments[0]=min(.1667,.16/min(w,h))  # same corner radius on every card size
    if line:b.line.color.rgb=rgb(line);b.line.width=Emu(10160)
    else:b.line.fill.background()
    return b
def card(s,x,y,w,h,accent,heading,body,hsize=16,bsize=13):
    box(s,x,y,w,h);box(s,x,y,0.07,h,fill=accent,line=None,shape=MSO_SHAPE.RECTANGLE)
    text(s,x+.2,y+.12,w-.32,.4,heading,hsize,TXT.get(accent,accent),True);text(s,x+.2,y+.55,w-.32,h-.62,body,bsize,INK)
def stat(s,x,y,w,value,label,color,size=32):
    text(s,x,y,w,.62,value,size,TXT.get(color,color),True,'c');text(s,x,y+.66,w,.5,label,12,CAPTION,False,'c')
def x_accuracy_axis(ch):
    """Numeric accuracy axis for XY charts: 20%-100% in steps of 20, titled."""
    xa=ch.category_axis;xa.minimum_scale=.2;xa.maximum_scale=1.05;xa.major_unit=.2;xa.has_major_gridlines=False
    xa.tick_labels.number_format='0%';xa.tick_labels.number_format_is_linked=False
    xa.has_title=True;tf=xa.axis_title.text_frame;tf.text='Controlled prediction accuracy';style_run(tf.paragraphs[0].runs[0],12,CAPTION,False)
def no_invert(series):
    """PowerPoint draws negative bars hollow unless every data point also opts out of inversion."""
    for dPt in series._element.findall(qn('c:dPt')):
        if dPt.find(qn('c:invertIfNegative')) is None:dPt.find(qn('c:idx')).addnext(etree.SubElement(dPt,qn('c:invertIfNegative'),val='0'))
def caption(s,t,y=6.38):text(s,.75,y,11.85,.32,t,12,CAPTION,True,'c')
def takeaway(s,t,y=6.0):text(s,.75,y,11.85,.42,t,16,INK,True,'c')
def arrow(s,x1,y1,x2,y2,color=MUTED,width=1.75,head=True):
    c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));c.line.color.rgb=rgb(color);c.line.width=Pt(width)
    if head:etree.SubElement(c.line._get_or_add_ln(),qn('a:tailEnd'),type='triangle',w='med',len='med')
    return c
def frame(prs,title,notes=''):
    s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb(BG)
    text(s,.55,.35,.55,.28,'00',12,ORANGE_TXT,True,'c','m');text(s,1.10,.28,11.55,.51,title,24,INK,True)
    text(s,.65,7.14,6.77,.21,FOOT,8,MUTED);text(s,12.15,7.12,.50,.21,'0',8,MUTED,False,'r')
    s.notes_slide.notes_text_frame.text=notes;return s
def picture(s,blob,x,y,w,h):
    """Insert an image inside the box (x, y, w, h) with its native aspect ratio, centred."""
    from PIL import Image
    iw,ih=Image.open(io.BytesIO(blob)).size;scale=min(w/iw,h/ih);pw,ph=iw*scale,ih*scale
    return s.shapes.add_picture(io.BytesIO(blob),Inches(x+(w-pw)/2),Inches(y+(h-ph)/2),Inches(pw),Inches(ph))
def style_chart(ch,title,legend=True,numfmt='+0.0;−0.0;0',size=12):
    ch.has_title=True;tf=ch.chart_title.text_frame;tf.text=title;style_run(tf.paragraphs[0].runs[0],15,INK,True)
    ch.font.size=Pt(size);ch.font.name=FONT;ch.font.color.rgb=rgb('4A5562')
    va=ch.value_axis;va.has_major_gridlines=True;va.major_gridlines.format.line.color.rgb=rgb('D9DFE5');va.major_gridlines.format.line.width=Pt(.75)
    va.format.line.color.rgb=rgb('D4D4D4');va.tick_labels.number_format=numfmt;va.tick_labels.number_format_is_linked=False
    ca=ch.category_axis;ca.format.line.color.rgb=rgb('9AA3AD')
    try:ca.tick_label_position=XL_TICK_LABEL_POSITION.LOW
    except Exception:pass
    if legend:ch.has_legend=True;ch.legend.position=XL_LEGEND_POSITION.BOTTOM;ch.legend.include_in_layout=False;ch.legend.font.size=Pt(size)
    else:ch.has_legend=False
    for el in (ch._chartSpace.chart.plotArea,ch._chartSpace):
        spPr=el.find(qn('c:spPr'))
        if spPr is None:spPr=etree.SubElement(el,qn('c:spPr'))
        for child in list(spPr):spPr.remove(child)
        etree.SubElement(etree.SubElement(spPr,qn('a:solidFill')),qn('a:srgbClr'),val='FFFFFF')
    cs=ch._chartSpace;spPr=cs.find(qn('c:spPr'));cs.remove(spPr);cs.find(qn('c:chart')).addnext(spPr)
def plot_layout(ch,x,y,w,h):
    """Fix the inner plot area (fractions of the chart) so overlay shapes can be placed exactly."""
    pa=ch._chartSpace.chart.plotArea;old=pa.find(qn('c:layout'))
    if old is not None:pa.remove(old)
    xml=('<c:layout xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"><c:manualLayout><c:layoutTarget val="inner"/><c:xMode val="edge"/><c:yMode val="edge"/>'
         f'<c:x val="{x}"/><c:y val="{y}"/><c:w val="{w}"/><c:h val="{h}"/></c:manualLayout></c:layout>')
    pa.insert(0,etree.fromstring(xml))
def err_bars(series,rows,color,direction=None):
    lit=lambda vals:'<c:numLit xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"><c:formatCode>General</c:formatCode><c:ptCount val="%d"/>%s</c:numLit>'%(len(vals),''.join(f'<c:pt idx="{i}"><c:v>{v:.6f}</c:v></c:pt>' for i,v in enumerate(vals)))
    xml=('<c:errBars xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
         +(f'<c:errDir val="{direction}"/>' if direction else '')+'<c:errBarType val="both"/><c:errValType val="cust"/><c:noEndCap val="0"/>'
         f'<c:plus>{lit([r["high"]-r["mean"] for r in rows])}</c:plus><c:minus>{lit([r["mean"]-r["low"] for r in rows])}</c:minus>'
         f'<c:spPr><a:ln w="15875"><a:solidFill><a:srgbClr val="{color}"/></a:solidFill></a:ln></c:spPr></c:errBars>')
    ser=series._element;anchor=ser.find(qn('c:xVal')) if ser.find(qn('c:xVal')) is not None else ser.find(qn('c:cat'))
    anchor.addprevious(etree.fromstring(xml))
def table(s,x,y,widths,rows,fills=None,size=13,hsize=12,row_h=.46,first_bold=True,align_first='l',aligns=None):
    t=s.shapes.add_table(len(rows),len(rows[0]),Inches(x),Inches(y),Inches(sum(widths)),Inches(row_h*len(rows))).table
    tblPr=t._tbl.tblPr;tblPr.set('bandRow','0');tblPr.set('firstRow','1')
    for j,w in enumerate(widths):t.columns[j].width=Inches(w)
    for i,row in enumerate(rows):
        t.rows[i].height=Inches(row_h)
        for j,val in enumerate(row):
            c=t.cell(i,j);c.fill.solid();c.fill.fore_color.rgb=rgb(INK if i==0 else (fills or {}).get((i,j),'FFFFFF'))
            c.margin_left=c.margin_right=Inches(.08);c.margin_top=c.margin_bottom=Inches(.03);c.vertical_anchor=MSO_ANCHOR.MIDDLE
            a=aligns[j] if aligns else (align_first if j==0 else 'c')
            p=c.text_frame.paragraphs[0];c.text_frame.word_wrap=True;p.alignment=PP_ALIGN.LEFT if a=='l' else PP_ALIGN.CENTER
            r=p.add_run();r.text=val;style_run(r,hsize if i==0 else size,'FFFFFF' if i==0 else INK,i==0 or (j==0 and first_bold))
    return t
def sig(r):return POS if r['low']>0 else NEG if r['high']<0 else None
def flow_box(s,x,y,w,h,label,color=INK,size=13):
    box(s,x,y,w,h,fill='FFFFFF',line='C9D1DA');text(s,x+.05,y,w-.1,h,label,size,color,False,'c','m')

# ---------- slides ----------
def s_title(prs,img):
    s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb(BG)
    box(s,7.45,0,5.885,7.5,fill='FFFFFF',line=None,shape=MSO_SHAPE.RECTANGLE);picture(s,img,7.6,.3,5.6,6.9)
    text(s,.75,1.05,6.4,1.1,['Revisiting Lifetime-Aware','VM Allocation'],30,INK,True)
    text(s,.75,2.55,6.3,.8,['Reconstruction, Diagnosis and','Prior-Fused Lifetime Prediction'],19,NAVY,True)
    text(s,.75,3.62,6.1,.32,'Lifetime prediction, confidence routing and prior fusion',13,CAPTION)
    for i,(label,color) in enumerate([('Synthetic reconstruction',NAVY),(f'{N_TEST} fresh test seeds',GREEN),('Pre-registered, Holm-corrected',ORANGE_TXT)]):
        x,y=.75+(i%2)*2.95,4.45+(i//2)*.55;box(s,x,y,2.75,.4,fill=color,line=None);text(s,x,y,2.75,.4,label,11,'FFFFFF',True,'c','m')
    text(s,.75,6.25,5,.3,'CE7490 Course Project',12,CAPTION,True);text(s,.75,6.6,5,.28,f'Revised {REVISED}',11,CAPTION)
    return s
def s_overview(prs):
    s=frame(prs,'Three questions, three answers')
    lava=lambda a:T('lava','best_fit',a)['mean'];host=T('host_ca','best_fit','synthetic')
    cards=[(NAVY,'RQ1 · Prediction quality','Does LAVA’s benefit depend on how accurate lifetime predictions are?',f'{num(lava("0.25"),2)} → {num(lava("1.0"),2)}{NB}pp',
            'Yes. LAVA loses to Best-Fit at 25% accuracy and gains clearly at 100%.'),
           (ORANGE,'RQ2 · Confidence routing','Does routing low-confidence VMs away from LAVA help?',f'{num(host["mean"],3)}{NB}pp',
            'No. Host-level routing is below Best-Fit; three measured causes contribute.'),
           (GREEN,'RQ3 · Prior fusion','Can the lifetime prediction itself be made robust?',f'{HOLDS_LAVA} of 4',
            f'Yes. Beats LAVA in {HOLDS_LAVA} of 4 pre-registered tests and no longer loses at low accuracy.')]
    for i,(color,head,question,big,answer) in enumerate(cards):
        x=.75+i*4.0;box(s,x,1.3,3.8,3.45);box(s,x,1.3,0.07,3.45,fill=color,line=None,shape=MSO_SHAPE.RECTANGLE)
        text(s,x+.2,1.42,3.45,.4,head,16,TXT.get(color,color),True);text(s,x+.2,1.87,3.45,.8,question,13,CAPTION)
        text(s,x+.2,2.75,3.45,.7,big,30,TXT.get(color,color),True);text(s,x+.2,3.6,3.45,1.05,answer,14,INK)
    takeaway(s,'Lesson: uncertainty helped inside the lifetime estimate, not as a policy switch.',5.2)
    caption(s,f'Fresh test seeds 1000–1099 (n{NB}={NB}{N_TEST}); all values relative to same-seed Best-Fit.',5.8)
    return s
def s_placement(prs,img):
    s=frame(prs,'VM placement and resource fragmentation')
    box(s,.75,1.3,7.2,4.25);picture(s,img,.85,1.4,7.0,4.05)
    for i,(color,head,body) in enumerate([(NAVY,'Multidimensional fragmentation','CPU and memory can each be free but unusable together.'),
                                          (ORANGE,'Time is the hidden dimension','VMs that fit equally well can exit at very different times.'),
                                          (GREEN,'Operational consequence','Empty hosts support maintenance, larger VMs and energy management.')]):
        card(s,8.25,1.3+i*1.45,4.33,1.25,color,head,body)
    takeaway(s,'Objective: maximize the time-weighted share of empty hosts without rejecting any VM.',5.85)
    bf=statistics.mean(IDX[(k,'synthetic','best_fit')]['empty_pct'] for k in SEEDS)
    caption(s,f'Scale: about {bf:.0f}% of host-time is empty under every policy (low load, long drain tail); 1{NB}pp ≈ 1.8 of 180 hosts.',6.38)
    return s
def s_lava(prs):
    s=frame(prs,'LAVA prediction and host lifecycle')
    for i,(color,head,body) in enumerate([(NAVY,'Distributions','Estimate a remaining-lifetime distribution for each VM.'),(ORANGE,'Re-prediction','Condition on uptime to revise the expected remaining lifetime.'),
                                          (GREEN,'Host lifecycle','Empty, open and recycling states guide placement.'),(RED,'Temporal placement','Align departures and fill recycling hosts with shorter VMs.')]):
        card(s,.75+i*2.99,1.25,2.82,1.55,color,head,body)
    y=3.55
    for label,color,x in [('empty',NAVY,1.6),('open',ORANGE,5.6),('recycling',GREEN,9.6)]:
        box(s,x,y,2.1,.55,fill='FFFFFF',line=color);text(s,x,y,2.1,.55,label,18,TXT.get(color,color),True,'c','m')
    arrow(s,3.75,y+.275,5.55,y+.275);arrow(s,7.75,y+.275,9.55,y+.275)
    text(s,3.75,y-.38,1.8,.34,'first VM',12,CAPTION,False,'c');text(s,7.55,y-.38,2.2,.34,'≥90% CPU or memory',12,CAPTION,False,'c')
    arrow(s,10.65,y+.6,10.65,y+.95,head=False);arrow(s,10.65,y+.95,2.65,y+.95,head=False);arrow(s,2.65,y+.95,2.65,y+.6)
    text(s,4.6,y+.97,4.2,.32,'last VM exits → host is empty again',12,CAPTION,False,'c')
    card(s,1.6,5.05,10.1,1.0,ORANGE,'Adaptation to mistakes','Residual cohorts lower the host class when they leave; missed deadlines raise it when VMs outlive their prediction.',hsize=15,bsize=13)
    caption(s,'Ling et al., “LAVA: Lifetime-Aware VM Allocation with Learned Distributions and Adaptation to Mispredictions”, MLSys 2025 (arXiv:2412.09840).',6.38)
    return s
def s_scope(prs):
    s=frame(prs,'Reconstruction scope and seed splits')
    for i,(color,head,body) in enumerate([(GREEN,'Recovered settings','180 hosts, 2,500 VMs per trace, 8 arrivals per hour, seed splits and metric.'),
                                          (RED,'Inferred defaults','Mixture parameters, VM shapes, scoring details and the controlled confidence model.'),
                                          (NAVY,'Unrecoverable inputs','Original traces, random streams, code history and per-seed results.')]):
        card(s,.75+i*4.0,1.25,3.8,1.45,color,head,body)
    rows=[['Split','Seeds','Used for'],['Initial sweep','0–19','RQ1 accuracy sweep; development of the new method'],['Validation','200–205','Report’s threshold selection; τ = 0.3 kept frozen'],
          ['Held-out','20–39','Report’s test split; rerun with all ten policies'],['Fresh test','1000–1099',f'Pre-registered evaluation of prior fusion (n{NB}={NB}{N_TEST})'],
          ['Robustness','2000–2049','Predictor variants, same-seed control and load × 2 (n = 50)'],['Threshold sensitivity','200–219','Descriptive τ scan for the new method (n = 20)']]
    table(s,.75,2.95,[2.6,1.7,7.53],rows,size=13,row_h=.46,aligns='lcl')
    return s
def s_methods(prs):
    s=frame(prs,'Methods compared')
    items=[(NAVY,'Best-Fit','Resource-only baseline: minimize residual CPU and memory.'),(ORANGE,'LA-Binary','One-shot short/long label at arrival, then Best-Fit.'),
           (GREEN,'NILAS','Re-predicted temporal extension cost, then Best-Fit.'),(NAVY,'LAVA','Lifetime classes and host states; adapts when VMs outlive their prediction.'),
           (RED,'VM-Level CA (VM-CA)','Per-arrival confidence gate between LAVA and Best-Fit on shared hosts.'),(GREEN,'Host-Level CA (Host-CA)','Same gate, but each host keeps one mode until it empties.')]
    for i,(color,head,body) in enumerate(items):card(s,.75+(i%3)*4.0,1.25+(i//3)*1.7,3.8,1.5,color,head,body)
    box(s,.75,4.75,11.83,1.3,fill='EEF6F0',line='C9DFCF');box(s,.75,4.75,0.07,1.3,fill=GREEN,line=None,shape=MSO_SHAPE.RECTANGLE)
    text(s,.95,4.87,11.4,.4,'New in this project: *-Bayes variants (RQ3)',16,GREEN,True)
    text(s,.95,5.3,11.4,.7,'NILAS-Bayes, LAVA-Bayes, VM-CA-Bayes and Host-CA-Bayes run exactly the same scheduling code as their base policy and only swap in the prior-fused lifetime predictor.',14,INK)
    return s
def s_settings(prs):
    s=frame(prs,'Prediction settings and information boundary')
    card(s,.75,1.25,3.85,1.35,GREEN,'Correct bucket',f'Log10 error with σ{NB}={NB}0.001.')
    card(s,.75,2.8,3.85,1.35,RED,'Incorrect bucket',f'Log10 error with σ{NB}={NB}3 (orders of magnitude).')
    card(s,.75,4.35,3.85,1.6,NAVY,'Information boundary','The scheduler never sees the bucket or the true lifetime.')
    text(s,4.95,1.2,7.6,.36,'Controlled accuracy = share of VMs in the correct bucket',13,CAPTION,True,'c')
    xs=[5.2,6.65,8.1,9.55,11.0]
    for i,(x,lvl) in enumerate(zip(xs,['25%','50%','75%','85%','100%'])):
        text(s,x,1.6,1.3,.45,lvl,20,NAVY,True,'c','m')
        if i<4:arrow(s,x+1.15,1.825,x+1.6,1.825,head=False)
    card(s,4.95,2.4,7.63,1.55,NAVY,'Default predictor','Profile-level two-exponential mixtures give the conditional remaining lifetime and a class probability. Numeric parameters are inferred.')
    card(s,4.95,4.15,3.7,1.8,ORANGE,'Oracle reference','True remaining lifetime; a full-information reference, not an upper bound.')
    card(s,8.88,4.15,3.7,1.8,GREEN,'Statistical protocol','Identical traces and noise for all policies; paired Student-t 95% intervals.')
    caption(s,'*-Bayes adds only the profile prior, which the default predictor already uses.',6.25)
    return s
def s_sweep(prs):
    s=frame(prs,'RQ1: Prediction accuracy drives LAVA’s gain')
    xs=[.25,.5,.75,.8,.85,.9,.95,1.];pts=[pick(SWEEP,'lava','best_fit',f'{x:g}' if x<1 else '1.0',split='initial') for x in xs]
    cd=XyChartData();sr=cd.add_series('LAVA − Best-Fit')
    for x,r in zip(xs,pts):sr.add_data_point(x,round(r['mean'],4))
    ch=s.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_LINES,Inches(.75),Inches(1.2),Inches(8.2),Inches(4.85),cd).chart
    style_chart(ch,'LAVA minus Best-Fit (pp), seeds 0–19',legend=False);ch.value_axis.minimum_scale=-.4;ch.value_axis.maximum_scale=.8;ch.value_axis.major_unit=.2
    x_accuracy_axis(ch)
    ser=ch.plots[0].series[0];ser.format.line.color.rgb=rgb(SLATE);ser.format.line.width=Pt(2.25);ser.smooth=False
    ser.marker.style=XL_MARKER_STYLE.CIRCLE;ser.marker.size=8;ser.marker.format.fill.solid();ser.marker.format.fill.fore_color.rgb=rgb(SLATE);ser.marker.format.line.color.rgb=rgb('FFFFFF')
    err_bars(ser,pts,SLATE,'y')
    d0=pick(SWEEP,'lava','best_fit','synthetic',split='initial');dh=pick(HELD,'lava','best_fit','synthetic',split='heldout');df=T('lava','best_fit','synthetic')
    card(s,9.25,1.25,3.33,1.4,NAVY,'First significant point',f'75% accuracy: {ci(pts[2])}{NB}pp',bsize=13)
    card(s,9.25,2.85,3.33,1.6,ORANGE,'Default predictor',f'{num(d0["mean"])}{NB}pp here; not replicated on held-out ({num(dh["mean"])}) or fresh seeds ({num(df["mean"])}).',bsize=13)
    card(s,9.25,4.65,3.33,1.4,GREEN,'Oracle reference',f'{ci(pick(SWEEP,"lava","best_fit","oracle",split="initial"))}{NB}pp',bsize=13)
    caption(s,'Initial split (n = 20 per point), paired 95% intervals. The 75% crossing point is specific to this reconstruction.',6.3)
    return s
def s_heldout(prs):
    s=frame(prs,'Held-out sensitivity and the RQ2 hypothesis')
    for i,(v,l,color) in enumerate([(f"{num(pick(HELD,'lava','best_fit','synthetic',split='heldout')['mean'])}{NB}pp",'LAVA, default predictor (held-out 20–39)',SLATE),
                                    (f"{num(pick(HELD,'lava','best_fit','oracle',split='heldout')['mean'])}{NB}pp",'LAVA with true lifetimes (held-out 20–39)',GREEN),
                                    (f"{num(T('lava','best_fit','0.25')['mean'])}{NB}pp",'LAVA at 25% accuracy (fresh seeds)',RED)]):
        box(s,.75+i*4.0,1.3,3.8,1.5);stat(s,.85+i*4.0,1.47,3.6,v,l,color)
    w,gap=3.6,.515
    for i,(color,head,body) in enumerate([(NAVY,'Observation','LAVA gains little with the default predictor and loses to Best-Fit when predictions are poor.'),
                                          (ORANGE,'Hypothesis','The earlier report’s idea: send low-confidence requests to Best-Fit, per VM or per host.'),
                                          (RED,'Constraint','Workload and threshold stay fixed; no search for a result that matches the earlier report.')]):
        x=.75+i*(w+gap);card(s,x,3.2,w,1.75,color,head,body,bsize=14)
        if i<2:arrow(s,x+w+.07,4.07,x+w+gap-.07,4.07)
    caption(s,'All values are paired differences relative to same-seed Best-Fit.',5.35)
    return s
def s_routing(prs):
    s=frame(prs,'RQ2: Two ways to route by confidence')
    for k,(title,color,yes,no,note) in enumerate([('VM-Level CA',NAVY,'LAVA scoring','Best-Fit scoring','Both decisions land on one shared set of hosts: modes mix on a host.'),
                                                   ('Host-Level CA',RED,'LAVA-mode hosts only','Best-Fit-mode hosts only','A host keeps its first mode until empty; cross-mode only if nothing else fits.')]):
        x0=.75+k*6.03;box(s,x0,1.2,5.8,3.55);text(s,x0+.2,1.28,5.4,.4,title,17,TXT.get(color,color),True)
        flow_box(s,x0+.25,2.35,1.25,.62,'Arrival');arrow(s,x0+1.55,2.66,x0+1.95,2.66)
        b=s.shapes.add_shape(MSO_SHAPE.DIAMOND,Inches(x0+2.0),Inches(2.0),Inches(1.55),Inches(1.32));b.fill.solid();b.fill.fore_color.rgb=rgb('FFF7EC');b.line.color.rgb=rgb(ORANGE);b.line.width=Emu(12700)
        text(s,x0+2.05,2.3,1.45,.72,['Confidence',f'≥{NB}τ{NB}?'],12,INK,True,'c','m')
        flow_box(s,x0+4.05,1.82,1.55,.62,yes,size=12);flow_box(s,x0+4.05,2.88,1.55,.62,no,size=12)
        arrow(s,x0+3.55,2.5,x0+4.0,2.13);arrow(s,x0+3.55,2.82,x0+4.0,3.19)
        text(s,x0+3.5,1.86,.6,.3,'yes',11,CAPTION,False,'c');text(s,x0+3.5,3.2,.6,.3,'no',11,CAPTION,False,'c')
        text(s,x0+.25,3.7,5.3,.9,note,13,INK)
    stats=[(f"{held_mean('vm_ca','cross_baseline_into_lava'):.0f} + {held_mean('vm_ca','cross_lava_into_baseline'):.0f} → 0",['cross-mode placements per trace,','VM-CA → Host-CA'],RED),
           (f"{100*held_mean('vm_ca','high_gate')/C['vms']:.0f}%",['of arrivals pass the gate,','identical in both variants'],NAVY),(f"{held_mean('host_ca','capacity_fallback'):.0f}",'capacity fallbacks in any run',GREEN)]
    for i,(v,l,color) in enumerate(stats):stat(s,.75+i*4.0,4.95,3.8,v,l,color,size=26)
    caption(s,'Mechanism check on held-out seeds 20–39, default predictor: isolation does exactly what it was designed to do.',6.4)
    return s
def s_negative(prs):
    s=frame(prs,'Diagnostic negative result: routing does not beat Best-Fit')
    items=[('LAVA','lava',SLATE),('VM-Level CA','vm_ca',NAVY),('Host-Level CA','host_ca',RED)];rows=[T(p,'best_fit','synthetic') for _,p,_ in items]
    cd=CategoryChartData();cd.categories=[f'{n}\n{num(r["mean"])}' for (n,_,_),r in zip(items,rows)];cd.add_series('vs Best-Fit',[round(r['mean'],4) for r in rows])
    ch=s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Inches(.75),Inches(1.2),Inches(7.3),Inches(4.85),cd).chart
    style_chart(ch,'Default predictor, fresh seeds: gain over Best-Fit (pp)',legend=False,numfmt='+0.00;−0.00;0');pl=ch.plots[0];pl.gap_width=90
    ch.value_axis.minimum_scale=-.2;ch.value_axis.maximum_scale=.1;ch.value_axis.major_unit=.05
    ser=pl.series[0];ser.invert_if_negative=False
    for i,(_,_,color) in enumerate(items):pt=ser.points[i];pt.format.fill.solid();pt.format.fill.fore_color.rgb=rgb(color)
    no_invert(ser);err_bars(ser,rows,INK)
    host,hold_bf,hold_vm,vm_hold=T('host_ca','best_fit','synthetic'),pick(HELD,'host_ca','best_fit','synthetic',split='heldout'),pick(HELD,'host_ca','vm_ca','synthetic',split='heldout'),pick(HELD,'vm_ca','best_fit','synthetic',split='heldout')
    h=HIST[('synthetic','host_ca')]
    card(s,8.3,1.25,4.28,1.5,RED,f'Fresh seeds (n{NB}={NB}{N_TEST})',f'Host-CA − Best-Fit: {ci(host)}{NB}pp, significantly below zero.',bsize=13)
    card(s,8.3,2.95,4.28,1.5,NAVY,'Held-out seeds 20–39',f'Host-CA {num(hold_bf["mean"])} vs Best-Fit, {num(hold_vm["mean"])} vs VM-CA; VM-CA’s {num(vm_hold["mean"])} did not replicate.',bsize=13)
    card(s,8.3,4.65,4.28,1.4,ORANGE,'Surviving report (historical)',f'Reported Host-CA − Best-Fit: {num(float(h["mean_pp"]))} [{num(float(h["low_pp"]))}, {num(float(h["high_pp"]))}]{NB}pp; not reproduced here.',bsize=13)
    caption(s,'Paired 95% intervals; threshold τ = 0.3 as in the report.',6.32)
    return s
def s_diagnosis(prs):
    s=frame(prs,'Why confidence routing does not help here')
    names=list(C['profiles']);cd=CategoryChartData();cd.categories=[n.capitalize() for n in names]
    cd.add_series('Probability of the most likely class',[round(max(profile(n)[0]),3) for n in names]);cd.add_series('Confidence used by the gate',[round(profile(n)[0][profile(n)[2]],3) for n in names])
    gx,gy,gw,gh=.75,1.2,7.3,4.85;ch=s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Inches(gx),Inches(gy),Inches(gw),Inches(gh),cd).chart
    style_chart(ch,'Predictability vs gate confidence at arrival',numfmt='0.0');pl=ch.plots[0];pl.gap_width=70;pl.overlap=-8
    ch.value_axis.maximum_scale=1.0;ch.value_axis.minimum_scale=0;ch.value_axis.major_unit=.2
    pl.has_data_labels=True;dl=pl.data_labels;dl.number_format='0.000';dl.number_format_is_linked=False;dl.position=XL_LABEL_POSITION.OUTSIDE_END;dl.font.size=Pt(12);dl.font.color.rgb=rgb(INK)
    for ser,color in zip(pl.series,[NAVY,ORANGE]):ser.format.fill.solid();ser.format.fill.fore_color.rgb=rgb(color)
    # Labels that would sit on the threshold line go inside their bar instead. A point-level label creates a
    # series-level label block that overrides the plot-level one, so restate the series defaults first.
    sl=pl.series[1].data_labels;sl.show_value=True;sl.number_format='0.000';sl.number_format_is_linked=False;sl.position=XL_LABEL_POSITION.OUTSIDE_END;sl.font.size=Pt(12);sl.font.color.rgb=rgb(INK)
    for j,n in enumerate(names):
        v=profile(n)[0][profile(n)[2]]
        if C['threshold']-.08<v<C['threshold']:dl2=pl.series[1].points[j].data_label;dl2.position=XL_LABEL_POSITION.INSIDE_END;dl2.font.size=Pt(12);dl2.font.color.rgb=rgb(INK)
    px,py,pw,ph=.09,.13,.79,.62;plot_layout(ch,px,py,pw,ph);ytau=gy+gh*(py+ph*(1-C['threshold']))
    ln=arrow(s,gx+gw*px,ytau,gx+gw*(px+pw),ytau,color=RED,width=1.5,head=False);ln.line.dash_style=7
    text(s,gx+gw*(px+pw)+.04,ytau-.16,.85,.32,f'τ{NB}={NB}{C["threshold"]}',12,RED,True,'l','m')
    win=T('host_ca','best_fit','synthetic','empty_pct_arrivals');i_pr,_,i_k=profile('interactive');s_pr,_,s_k=profile('service')
    card(s,8.3,1.25,4.28,1.5,NAVY,'Inverted confidence',f'Interactive VMs are {i_pr[0]:.0%} likely to exit within 1{NB}h but get {i_pr[i_k]:.3f}; service VMs get {s_pr[s_k]:.3f}.',bsize=13)
    card(s,8.3,2.95,4.28,1.5,ORANGE,'Blind gate',f'Under controlled noise ≥{min(gate_rate(a,"vm_ca") for a in ACC):.2f}% of arrivals pass τ{NB}={NB}0.3, so both variants essentially reduce to LAVA.',bsize=13)
    card(s,8.3,4.65,4.28,1.4,RED,'Isolation costs packing',f'Host-CA loses {abs(win["mean"]):.2f}{NB}pp of empty hosts during arrivals; the tail recovers about {tail_recovery():.0%}.',bsize=13)
    caption(s,'Default predictor at arrival; gate rates and windows on fresh seeds. Fixing the inversion alone did not help on development seeds either.',6.32)
    return s
def s_method(prs):
    s=frame(prs,'RQ3: Prior-fused lifetime prediction')
    steps=[(NAVY,'Inputs','Profile prior, noisy value, accuracy and σ'),(ORANGE,'Posterior over T',f'Log-grid atoms; values at the 336{NB}h cap are censored'),
           (GREEN,'Uptime filter','Keep only atoms above the current age'),(NAVY,'Outputs','Mean remaining lifetime and class confidence'),(RED,'Scheduler','NILAS, LAVA and CA rules unchanged')]
    w,gap=2.03,.42
    for i,(color,head,body) in enumerate(steps):
        x=.75+i*(w+gap);box(s,x,1.25,w,1.6);text(s,x+.12,1.35,w-.24,.38,head,15,TXT.get(color,color),True);text(s,x+.12,1.78,w-.24,1.0,body,13,INK)
        if i<len(steps)-1:arrow(s,x+w+.05,2.05,x+w+gap-.05,2.05)
    f1=predict_pair('interactive','0.25',300.)[0][0];c300,r300=predict_pair('service','1.0',float(C['cap_hours']),300.)
    prior=mixture_remaining(C['profiles']['interactive'],0)[0]
    card(s,.75,3.15,3.8,1.75,NAVY,'Implausible values fall back',f'Interactive VM at 25% accuracy predicted 300{NB}h: the original predictor uses 300{NB}h, the fused one {f1:.2f}{NB}h (prior mean {prior:.2f}{NB}h).')
    card(s,4.77,3.15,3.8,1.75,ORANGE,'Censoring at the cap',f'A value at 336{NB}h means “at least 336{NB}h”. At uptime 300{NB}h the original expects {r300[0]:.0f}{NB}h left, the fused one {c300[0]:.0f}{NB}h.')
    card(s,8.78,3.15,3.8,1.75,GREEN,'No overdue jump',f'A VM that outlives its prediction keeps a posterior estimate instead of jumping to “remaining{NB}={NB}uptime”.')
    takeaway(s,'Same information as the original confidence model plus the profile prior; never the bucket or the true lifetime.',5.25)
    n=len((root/'lava/model.py').read_text().split('class BayesPredictor',1)[1].splitlines())
    caption(s,f'Implementation: BayesPredictor in lava/model.py ({n} lines), unit-tested in tests/test_bayes.py; examples computed by the code.',5.8)
    return s
def s_protocol(prs):
    s=frame(prs,'Pre-registered evaluation protocol')
    reg=PROTO['registered_at'];stamp=f"{int(reg[8:10])} Oct, {reg[11:16]} (UTC+8)"
    steps=[(NAVY,'1  Develop','Seeds 0–19 only. Two other ideas rejected there.'),(ORANGE,'2  Register',f'{stamp}: hypotheses, seeds, τ, decision rule and code hashes frozen.'),
           (GREEN,'3  Test',f'Fresh seeds 1000–1099 (n{NB}={NB}{N_TEST}); τ{NB}={NB}0.3 from the report, nothing tuned.'),(RED,'4  Infer',['Eight paired t-tests,',f'Holm–Bonferroni at α{NB}={NB}0.05.'])]
    w,gap=2.6,.477
    for i,(color,head,body) in enumerate(steps):
        x=.75+i*(w+gap);card(s,x,1.25,w,1.85,color,head,body,hsize=16,bsize=13)
        if i<3:arrow(s,x+w+.06,2.17,x+w+gap-.06,2.17)
    box(s,.75,3.4,11.83,1.75);text(s,.95,3.48,3,.36,'Compute',15,INK,True)
    workers=sorted({MAN[k]['workers'] for k in SUITES});assert len(workers)==1
    fails=sum(int(r['failures']) for k in SUITES for r in csv.DictReader((RES/f'improvement/{k}/raw.csv').open()))
    for i,(v,l) in enumerate([(f'{IMP_RUNS:,}','simulations'),(f'{workers[0]}',['low-priority CPU processes','on the H20 server']),(f'≈{IMP_MIN:.0f}{NB}min','wall time; no GPU needed'),(f'{fails}','scheduling failures')]):
        stat(s,.85+i*2.95,3.8,2.85,v,l,INK,size=26)
    caption(s,'Also run: robustness seeds 2000–2049 (control and load × 2 added after review), τ scan on 200–219, and the report’s held-out seeds 20–39.',5.5)
    caption(s,'Only the eight registered tests are multiplicity-adjusted; everything else is descriptive.',5.95)
    return s
def s_results(prs):
    s=frame(prs,'Fusion removes LAVA’s low-accuracy penalty')
    xs=[.25,.5,.75,.85,1.];series=[('VM-CA-Bayes (registered)','vm_ca_bayes',NAVY,XL_MARKER_STYLE.CIRCLE,-.008),('LAVA-Bayes (reference)','lava_bayes',ORANGE,XL_MARKER_STYLE.SQUARE,.008),('LAVA','lava',SLATE,XL_MARKER_STYLE.DIAMOND,0.)]
    rows={p:[T(p,'best_fit',a) for a in ACC] for _,p,_,_,_ in series};cd=XyChartData()
    for name,p,_,_,off in series:
        sr=cd.add_series(name)
        for x,r in zip(xs,rows[p]):sr.add_data_point(round(x+off,4),round(r['mean'],4))
    ch=s.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_LINES,Inches(.75),Inches(1.2),Inches(8.2),Inches(4.85),cd).chart
    style_chart(ch,'Gain over Best-Fit (pp), fresh seeds 1000–1099');ch.value_axis.minimum_scale=-.3;ch.value_axis.maximum_scale=.7;ch.value_axis.major_unit=.1
    x_accuracy_axis(ch)
    for ser,(name,p,color,mk,_) in zip(ch.plots[0].series,series):
        ser.format.line.color.rgb=rgb(color);ser.format.line.width=Pt(2.25);ser.smooth=False
        ser.marker.style=mk;ser.marker.size=7;ser.marker.format.fill.solid();ser.marker.format.fill.fore_color.rgb=rgb(color);ser.marker.format.line.color.rgb=rgb('FFFFFF')
        err_bars(ser,rows[p],color,'y')
    fam=[FAMILY[(a,'lava')] for a in REG];b25=FAMILY[('0.25','best_fit')];oracle=T('lava','best_fit','oracle')
    card(s,9.25,1.25,3.33,1.55,NAVY,'Registered: vs LAVA',f"{', '.join(num(r['mean']) for r in fam)}{NB}pp at 25–85%; {HOLDS_LAVA} of 4 hold.",bsize=13)
    card(s,9.25,2.98,3.33,1.55,GREEN,'Registered: vs Best-Fit',f"Holds at 50–85%; at 25%: {num(b25['mean'])}{NB}pp, not significant.",bsize=13)
    card(s,9.25,4.71,3.33,1.34,ORANGE,'Oracle reference',f"LAVA with true lifetimes: {num(oracle['mean'])}{NB}pp",bsize=13)
    caption(s,'Whiskers: unadjusted 95% intervals (n = 100). LAVA-Bayes was not registered and is shown for reference.',6.3)
    return s
def s_table(prs):
    s=frame(prs,f'Pre-registered results: {HOLDS_LAVA} of 4 hold against LAVA')
    rows=[['Accuracy','VM-CA-Bayes − LAVA','Holm p','Holds','VM-CA-Bayes − Best-Fit','Holm p','Holds']];fills={}
    for i,a in enumerate(REG,1):
        l,b=FAMILY[(a,'lava')],FAMILY[(a,'best_fit')];rows.append([lab(a),ci(l),pf(l['p_holm']),'yes' if ok(l) else 'no',ci(b),pf(b['p_holm']),'yes' if ok(b) else 'no'])
        for j,r in [(1,l),(4,b)]:
            if ok(r):fills[(i,j)]=POS
    table(s,.75,1.25,[1.39,3.05,1.15,1.0,3.05,1.15,1.04],rows,fills,size=13,row_h=.5)
    l75,b25=FAMILY[('0.75','lava')],FAMILY[('0.25','best_fit')];lb,nb=T('lava_bayes','best_fit','0.25'),T('nilas_bayes','best_fit','0.25')
    card(s,.75,4.0,5.8,1.35,ORANGE,'Not supported vs LAVA at 75%',f'{ci(l75)}{NB}pp: the interval includes zero (Holm p{NB}={NB}{pf(l75["p_holm"])}).',bsize=13)
    card(s,6.78,4.0,5.8,1.35,RED,'Not supported vs Best-Fit at 25%',f'{ci(b25)}{NB}pp. Secondary, unadjusted: NILAS-Bayes {num(nb["mean"])} (p{NB}{"<" if nb["p"]<1e-3 else "="}{NB}{pf(nb["p"]).lstrip("<")}), LAVA-Bayes {num(lb["mean"])} (p{NB}={NB}{pf(lb["p"])}).',bsize=13)
    caption(s,'Decision rule fixed before testing: mean > 0 and Holm-adjusted p < 0.05. Intervals are unadjusted 95% CIs.',5.65)
    caption(s,'Plain LAVA would also pass the Best-Fit tests at 75% and 85%, so those two say little about fusion itself.',6.05)
    return s
def s_decomp(prs):
    s=frame(prs,'Post hoc: the gain comes from the estimate, not the gate')
    noisy=[T('lava_bayes','lava',a) for a in REG];prior=[cross('lava_bayes',a,'lava','synthetic') for a in REG];gate=[posthoc('vm_ca_bayes','lava_bayes',a) for a in REG]
    cd=CategoryChartData();cd.categories=[lab(a) for a in REG]
    for name,rs in [('Fusion vs noisy values alone',noisy),('Fusion vs prior alone',prior),('Gate on top of fusion',gate)]:cd.add_series(name,[round(r['mean'],4) for r in rs])
    ch=s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Inches(.75),Inches(1.2),Inches(7.3),Inches(4.85),cd).chart
    style_chart(ch,'Paired difference (pp) by controlled accuracy',numfmt='+0.0;−0.0;0');pl=ch.plots[0];pl.gap_width=70;pl.overlap=-6
    ch.value_axis.minimum_scale=-.1;ch.value_axis.maximum_scale=.6;ch.value_axis.major_unit=.1
    for ser,color,rs in zip(pl.series,[NAVY,GREEN,ORANGE],[noisy,prior,gate]):
        ser.invert_if_negative=False;ser.format.fill.solid();ser.format.fill.fore_color.rgb=rgb(color);err_bars(ser,rs,INK)
    fc=[100*GB[(a,'fused')]['correct_pass'] for a in REG];fi=[100*GB[(a,'fused')]['incorrect_pass'] for a in REG];hw=max((r['high']-r['low'])/2 for r in gate)
    ns_noisy=[lab(a) for a,r in zip(REG,noisy) if r['low']<=0];assert all(r['low']>0 for r in prior)
    card(s,8.3,1.25,4.28,1.5,GREEN,'Fusion vs its ingredients',f"{num(min(r['mean'] for r in prior),2)} to {num(max(r['mean'] for r in prior),2)}{NB}pp over the prior alone; {num(min(r['mean'] for r in noisy),2)} to {num(max(r['mean'] for r in noisy),2)} over noisy values alone"+(f" (n.s. at {', '.join(ns_noisy)})." if ns_noisy else '.'),bsize=13)
    card(s,8.3,2.95,4.28,1.5,ORANGE,'Gate on top',f"{num(min(r['mean'] for r in gate),2)} to {num(max(r['mean'] for r in gate),2)}{NB}pp; every 95% interval includes zero (half-width ≈{NB}{hw:.2f}). The development-seed gain did not replicate.",bsize=13)
    card(s,8.3,4.65,4.28,1.4,NAVY,'Weakly selective gate',f'Routes {min(fc):.0f}–{max(fc):.0f}% of correct and {min(fi):.0f}–{max(fi):.0f}% of wrong predictions to LAVA.',bsize=13)
    caption(s,'Post hoc, not registered; fresh seeds (n = 100). “Prior alone” = the same seeds under the default predictor.',6.32)
    return s
def s_robust(prs):
    s=frame(prs,'Robustness and the failure mode')
    rlabels=[('protocol','Protocol, same seeds'),('bad_sigma=1','Incorrect bucket σ = 1'),('bad_sigma=0.5','Incorrect bucket σ = 0.5'),('fusion_assumes_3_true_1','Fusion assumes σ = 3 (true 1)'),
             ('fusion_assumes_1_true_3','Fusion assumes σ = 1 (true 3)'),('prior_means_x0.5','Prior means × 0.5'),('prior_means_x2','Prior means × 2'),('load_x2','Arrival rate × 2')]
    rows=[['VM-CA-Bayes − LAVA (pp)']+[lab(a) for a in REG]];fills={}
    for i,(v,name) in enumerate(rlabels,1):
        rs=[pick(IMP['robust_control'],'vm_ca_bayes','lava',a) if v=='protocol' else pick(IMP['robust_load'],'vm_ca_bayes','lava',a) if v=='load_x2' else pick(IMP['robust'],'vm_ca_bayes','lava',a,variant=v) for a in REG]
        rows.append([name]+[num(r['mean']) for r in rs]);fills.update({(i,j):sig(r) for j,r in enumerate(rs,1) if sig(r)})
    table(s,.75,1.2,[3.75,1.15,1.15,1.15,1.15],rows,fills,size=13,row_h=.5)
    x05=pick(IMP['robust'],'vm_ca_bayes','lava','0.85',variant='prior_means_x0.5');lo=[pick(IMP['robust_load'],'vm_ca_bayes','lava',a)['mean'] for a in REG[:3]]
    cells=[c for row in rows[1:] for c in row[1:]];neg=sum(1 for v in fills.values() if v==NEG);assert neg==1 and x05['high']<0  # text below names that one cell
    card(s,9.45,1.2,3.13,1.5,RED,'Failure mode: short prior',f'Prior means × 0.5 at 85% is the only significantly negative cell ({num(x05["mean"])}{NB}pp; {neg} of {len(cells)}).',bsize=12.5)
    card(s,9.45,2.87,3.13,1.5,NAVY,f'Load × 2 ({2*C["arrival_rate"]} VMs/h)',f"Gain over LAVA rises to {', '.join(num(v) for v in lo)}{NB}pp at 25%, 50% and 75%.",bsize=12.5)
    card(s,9.45,4.54,3.13,1.5,ORANGE,'Threshold τ',f'0.1–0.5: positive mean at every accuracy; τ{NB}≥{NB}0.7 loses the gain at 50–85% (backup slide).',bsize=12.5)
    caption(s,'Seeds 2000–2049 (n = 50), descriptive. Green / red: 95% interval above / below zero.',6.2)
    caption(s,'Same-seed control and load × 2 were added after review; the registered 90-host variant only rescales the 180-host run and is omitted.',6.55)
    return s
def s_findings(prs):
    s=frame(prs,'Findings and limits')
    text(s,.75,1.15,6.9,.45,'What we found',20,NAVY,True)
    lava=lambda a:T('lava','best_fit',a)['mean']
    for i,(color,head,body) in enumerate([(NAVY,'RQ1: prediction quality is first-order',f'LAVA moves from {num(lava("0.25"))} to {num(lava("1.0"))}{NB}pp between 25% and 100% accuracy.'),
                                          (ORANGE,'RQ2: a diagnostic negative result','Inverted confidence, a blind gate and the cost of isolation each contribute to the missing gain.'),
                                          (GREEN,'RQ3: prior fusion works',f'Beats LAVA in {HOLDS_LAVA} of 4 registered tests; post hoc, the estimate rather than the gate appears to carry the gain.')]):
        card(s,.75,1.7+i*1.42,6.9,1.27,color,head,body,hsize=16,bsize=13)
    box(s,7.95,1.15,4.63,4.52);text(s,8.15,1.25,4.3,.42,'Limits',20,RED,True)
    text(s,8.15,1.8,4.3,3.8,[[('Synthetic workload: ',{'bold':True}),('about 20% offered load; the drain tail is about two thirds of the metric.',{})],
                              [('Calibrated prior: ',{'bold':True}),('the generating distribution; halved means fail at 85%.',{})],
                              [('Reconstruction: ',{'bold':True}),('NILAS beats LAVA at most accuracies here; original code and traces were lost.',{})],
                              [('Statistics: ',{'bold':True}),('registration is self-timestamped; robustness cells are descriptive.',{})]],14,INK,space=14)
    takeaway(s,'Lesson: uncertainty helped inside the lifetime estimate, not as a policy switch.',6.0)
    return s
def s_close(prs):
    s=frame(prs,'Thank you')
    text(s,.75,1.6,11.83,.9,'Questions?',40,INK,True,'c')
    lava=lambda a:T('lava','best_fit',a)['mean']
    for i,(v,l,color) in enumerate([(f'{num(lava("0.25"),2)} → {num(lava("1.0"),2)}{NB}pp','RQ1: LAVA vs Best-Fit, 25% → 100% accuracy',NAVY),
                                    (f"{num(T('host_ca','best_fit','synthetic')['mean'])}{NB}pp",'RQ2: Host-Level CA vs Best-Fit',ORANGE),(f'{HOLDS_LAVA} of 4','RQ3: registered wins against LAVA',GREEN)]):
        box(s,.75+i*4.0,3.0,3.8,1.45);stat(s,.85+i*4.0,3.2,3.6,v,l,color,size=26)
    caption(s,'Report: Revisiting_Lifetime_Aware_VM_Allocation_Report_v3.pdf · Artifact: lava_reproduction/ (code, raw results, registration).',4.95)
    return s
def s_backup_tau(prs):
    s=frame(prs,'Backup: threshold sensitivity')
    taus=MAN['tau']['thresholds'];rows=[['τ']+[lab(a) for a in REG]];fills={}
    for i,t in enumerate(taus,1):
        rs=[pick(IMP['tau'],'vm_ca_bayes','lava',a,threshold=t) for a in REG];rows.append([f'{t}']+[ci(r) for r in rs]);fills.update({(i,j):sig(r) for j,r in enumerate(rs,1) if sig(r)})
    table(s,.75,1.25,[1.2,2.65,2.65,2.65,2.65],rows,fills,size=13,row_h=.5,align_first='c')
    caption(s,'VM-CA-Bayes minus LAVA (pp), seeds 200–219 (n = 20), descriptive; τ = 0.3 was frozen from the report, not chosen here.',5.65)
    raw=list(csv.DictReader((RES/'improvement/tau/raw.csv').open()))
    share=lambda t,a:100*statistics.mean(int(r['high_gate'])/int(r['accepted']) for r in raw if float(r['threshold'])==t and r['setting']==a and r['policy']=='vm_ca_bayes')
    hi=[share(t,a) for t in taus if t>=.7 for a in ('0.75','0.85')];lo=[share(.3,a) for a in ('0.75','0.85')]
    caption(s,f'At τ ≥ 0.7 the gate still sends {min(hi):.0f}–{max(hi):.0f}% of arrivals to LAVA at 75–85% accuracy ({min(lo):.0f}–{max(lo):.0f}% at τ = 0.3).',6.1)
    return s
def s_backup_all(prs):
    s=frame(prs,'Backup: all policies versus Best-Fit (fresh seeds)')
    pols=['lava','nilas','vm_ca','host_ca','lava_bayes','nilas_bayes','vm_ca_bayes','host_ca_bayes'];names=['LAVA','NILAS','VM-CA','Host-CA','LAVA-B','NILAS-B','VM-CA-B','Host-CA-B']
    rows=[['Setting']+names];fills={}
    for i,a in enumerate(ACC+['synthetic','oracle'],1):
        rs=[T(p,'best_fit',a) for p in pols];rows.append([lab(a)]+[num(r['mean']) for r in rs]);fills.update({(i,j):sig(r) for j,r in enumerate(rs,1) if sig(r)})
    table(s,.75,1.25,[1.55]+[1.285]*8,rows,fills,size=12,hsize=11.5,row_h=.5)
    caption(s,'Mean pp versus same-seed Best-Fit (n = 100); “-B” = prior-fused variant. Green / red: 95% interval above / below zero.',5.65)
    caption(s,'Under the default predictor and the Oracle the fused variants equal their base policies by construction.',6.1)
    return s

def speaker_notes():
    sw=lambda st:pick(SWEEP,'lava','best_fit',st,split='initial');hd=lambda p,b,st:pick(HELD,p,b,st,split='heldout')
    i_pr,_,i_k=profile('interactive');s_pr,_,s_k=profile('service');win=T('host_ca','best_fit','synthetic','empty_pct_arrivals')
    gate=[posthoc('vm_ca_bayes','lava_bayes',a) for a in REG];prior=[cross('lava_bayes',a,'lava','synthetic') for a in REG]
    x05=pick(IMP['robust'],'vm_ca_bayes','lava','0.85',variant='prior_means_x0.5');f1=predict_pair('interactive','0.25',300.)[0][0]
    fam=[FAMILY[(a,'lava')] for a in REG];reg=PROTO['registered_at'][:16].replace('T',' ')
    return [
     "Talking points: Our project revisits LAVA, Google's lifetime-aware VM allocator from MLSys 2025. The original course code and traces were lost, so we rebuilt the system in a transparent simulator and used it to ask how lifetime-prediction quality and uncertainty handling change the scheduling benefit.\nTiming: about 18 minutes for the 20 main slides; slides 21-22 are backup.",
     f"Talking points: The whole talk on one slide. RQ1: prediction quality is first-order; LAVA loses to Best-Fit at 25% accuracy and gains clearly at 100%. RQ2: routing low-confidence VMs away from LAVA does not help; host-level routing is even below Best-Fit, and we measure three causes. RQ3: fusing each noisy prediction with the profile's lifetime distribution fixes the low-accuracy problem; it beats LAVA in {HOLDS_LAVA} of 4 pre-registered tests on {N_TEST} fresh seeds. New in this revision: the RQ2 diagnosis and the RQ3 method.\nSource: results/improvement/test/summary.json.",
     "Talking points: VM placement is multi-dimensional bin packing. CPU and memory can each be free but unusable together. Time is the hidden dimension: two VMs that fit equally well today may leave at very different times, so one long-lived VM can keep an otherwise empty host busy. Our metric is the time-weighted share of empty hosts. Because load is low and the drain tail is long, about 90% of host-time is empty under every policy, so differences are small in percentage points: 1 pp is about 1.8 hosts.\nSource: report Sections 1 and 5.",
     "Talking points: LAVA predicts a lifetime distribution, re-predicts the remaining lifetime as a VM ages, and manages each host as empty, open or recycling. Hosts carry a lifetime class with boundaries at 1, 10 and 100 hours. When the residual cohort leaves early the host moves down a class; when VMs outlive a deadline it moves up. This is how LAVA adapts to mispredictions.\nSource: Ling et al., arXiv:2412.09840, Sections 4.2-4.3.",
     "Talking points: The original code and traces were lost, so we rebuilt the simulator from the surviving report: 180 hosts, 2,500 VMs per trace and 8 arrivals per hour. Values we could not recover are declared once as inferred defaults and never tuned to match old numbers. The table lists every seed split; the new method is judged only on the fresh test seeds 1000-1099.\nSource: REPRODUCTION_DIFFERENCES.md; configs/provenance.json.",
     "Talking points: Best-Fit only packs resources. LA-Binary and NILAS add simple lifetime signals; LAVA adds classes and host states. VM-Level CA and Host-Level CA route low-confidence requests to Best-Fit, per VM or per host. Each *-Bayes variant runs exactly the same scheduling code and only swaps in the prior-fused predictor, so any change comes from the prediction.\nSource: lava/simulator.py.",
     "Talking points: Controlled accuracy follows LAVA's Appendix G.1: with probability equal to the accuracy the log10 error has sigma 0.001, otherwise sigma 3, i.e. errors of several orders of magnitude. The default predictor uses profile-level mixtures; the Oracle knows the true remaining lifetime and is a reference, not an upper bound. The scheduler never sees the bucket or the true lifetime; the fused predictor additionally uses the profile prior that the default predictor already has.\nSource: lava/model.py.",
     f"Talking points: RQ1. LAVA's gain over Best-Fit grows with accuracy; 75% is the first tested level whose interval excludes zero ({ci(sw('0.75'))} pp), and the Oracle reaches {ci(sw('oracle'))} pp. Under the default predictor LAVA's small gain here did not replicate on later seeds. The accuracy axis is to scale.\nSource: results/sweep/summary.json.",
     f"Talking points: On held-out seeds the default predictor gives LAVA essentially nothing ({num(hd('lava','best_fit','synthetic')['mean'])} pp) while the Oracle gives {num(hd('lava','best_fit','oracle')['mean'])} pp, and on fresh seeds LAVA at 25% accuracy is {num(T('lava','best_fit','0.25')['mean'])} pp. The earlier report's answer was confidence routing; we test it with the workload and threshold fixed.\nSource: results/heldout and results/improvement/test.",
     f"Talking points: Two designs. VM-Level CA gates each arrival between LAVA and Best-Fit on shared hosts, so decisions of both kinds mix on a host. Host-Level CA keeps each host in its first mode until it is empty. With the same gate sequence, isolation removes all cross-mode placements, about {held_mean('vm_ca','cross_baseline_into_lava'):.0f} plus {held_mean('vm_ca','cross_lava_into_baseline'):.0f} per trace, and never needs a capacity fallback. Both variants see the same gate decisions: about {100*held_mean('vm_ca','high_gate')/C['vms']:.0f}% of arrivals pass under the default predictor. So the mechanism works as designed.\nSource: results/heldout/raw.csv.",
     f"Talking points: But it does not help. On {N_TEST} fresh seeds Host-Level CA is significantly below Best-Fit ({ci(T('host_ca','best_fit','synthetic'))} pp). On the held-out seeds it was also below VM-CA, and VM-CA's small held-out gain did not replicate. The surviving report had claimed +0.241 pp; we could not reproduce it. We treat this as a diagnostic negative result.\nSource: results/improvement/test; results/heldout; reference/report_table1.csv.",
     f"Talking points: Three measured causes. One, confidence is inverted: interactive VMs are {i_pr[0]:.0%} likely to exit within an hour but get confidence {i_pr[i_k]:.3f}, because confidence measures the class containing the mean of a bimodal mixture; service VMs get {s_pr[s_k]:.3f}. Two, under controlled noise the gate passes almost every request. Three, isolation costs about {abs(win['mean']):.2f} pp of empty hosts during arrivals. Fixing the inversion alone did not help on development seeds, so these factors contribute; none is a full explanation on its own.\nSource: configs/report_scale.json; results/improvement/test; results/improvement/dev.",
     f"Talking points: So we repair the estimate instead of switching policies. Bayes' rule combines each noisy prediction with the profile's lifetime distribution; the 14-day cap is treated as censoring; the posterior is conditioned on uptime. The output interface is unchanged, so every scheduling rule stays the same. Example: an interactive VM predicted at 300 hours is implausible under its prior, so the fused estimate is {f1:.2f} hours. No extra information is used beyond the error model and the profile model the scheduler already had.\nSource: lava/model.py (BayesPredictor).",
     f"Talking points: To avoid choosing a method after seeing results, we developed on seeds 0-19 only, then froze hypotheses, test seeds, threshold, decision rule and code hashes at {reg} (UTC+8) before any test run. The registration is a local, self-timestamped file. The test uses {N_TEST} fresh seeds and eight paired t-tests with Holm correction. All {IMP_RUNS:,} improvement simulations ran in about {IMP_MIN:.0f} minutes on 64 CPU processes of our H20 server; no GPU is needed.\nSource: configs/improvement_protocol.json; results/improvement/*/manifest.json.",
     f"Talking points: Main result. The grey line is the original LAVA: below Best-Fit at 25% and about zero at 50%. The registered variant, VM-CA-Bayes, differs from LAVA by {', '.join(num(r['mean']) for r in fam)} pp at 25-85% accuracy; all but the 75% difference are significant after Holm correction. LAVA-Bayes, shown for reference, behaves almost identically. At 100% all methods approach the Oracle reference.\nSource: results/improvement/test/summary.json.",
     f"Talking points: {HOLDS_LAVA} of 4 registered tests against LAVA hold after Holm correction; at 75% the interval includes zero. Against Best-Fit 3 of 4 hold, but the 75% and 85% tests would also pass for plain LAVA, so the fusion-specific evidence is the LAVA tests and the 50% Best-Fit test. At 25% the Best-Fit test fails; the NILAS-Bayes and LAVA-Bayes results there are secondary and unadjusted.\nSource: results/improvement/test/summary.json (primary = true).",
     f"Talking points: Where does the gain come from? This is post hoc. Fusion improves on both ingredients at almost every accuracy: the prior alone (the same seeds under the default predictor, {num(min(r['mean'] for r in prior),2)} to {num(max(r['mean'] for r in prior),2)} pp) and the noisy predictions alone (not significant at 75%). Adding the gate on top changes little ({num(min(r['mean'] for r in gate))} to {num(max(r['mean'] for r in gate))} pp; every 95% interval includes zero, half-width about {max((r['high']-r['low'])/2 for r in gate):.2f}). The fused gate is only weakly selective, so the estimate rather than the gate appears to absorb most of the damage from wrong predictions. On development seeds the gate seemed to help; that did not replicate.\nSource: results/improvement/test/raw.csv; results/improvement/gate_by_bucket.json.",
     f"Talking points: Robustness on 50 fresh seeds, compared with the protocol on the same seeds. Milder errors and misspecified error spreads shrink the gain but never make it significantly negative. Doubling the arrival rate makes fusion more valuable at low accuracy. The failure mode is a prior that underestimates lifetimes: halved prior means are worse than LAVA at 85% ({ci(x05)} pp, one descriptive cell). Thresholds from 0.1 to 0.5 give positive mean gains; 0.3 was fixed in advance. Be open that the same-seed control and the load test were added after review: the registered 90-host variant only rescaled the 180-host run, so it is not shown.\nSource: results/improvement/robust, robust_control, robust_load, tau.",
     "Talking points: Summary. Prediction quality is first-order. Confidence routing is a diagnostic negative result with measurable causes. Prior fusion is a drop-in predictor change: it improves NILAS, LAVA and VM-CA at most accuracies but makes Host-CA worse, and 3 of 4 registered tests against LAVA hold. Limits: a synthetic workload whose metric is dominated by the drain tail, a calibrated prior, a reconstruction in which NILAS beats LAVA at most accuracies, and a self-timestamped registration.\nSource: report Sections 11-12.",
     "Talking points: Thank you. The three numbers summarize the answers to RQ1, RQ2 and RQ3. The report and the full artifact, with raw per-seed results and the registration file, are available for checking.",
     "Backup: threshold sensitivity on seeds 200-219. Use if asked whether tau = 0.3 was cherry-picked: it was frozen from the earlier report, and thresholds from 0.1 to 0.5 all give positive means.\nSource: results/improvement/tau.",
     "Backup: every policy on the fresh test seeds. Use if asked about NILAS (it beats LAVA at most accuracies here) or LA-Binary (slightly below Best-Fit everywhere).\nSource: results/improvement/test/summary.json."]

def renumber(prs):
    for i,s in enumerate(prs.slides,1):
        if i==1:continue
        for sh in s.shapes:
            if not sh.has_text_frame:continue
            t=sh.text_frame.text.strip();r=sh.text_frame.paragraphs[0].runs
            if t.isdigit() and abs(sh.left-Inches(.55))<Inches(.05) and abs(sh.top-Inches(.35))<Inches(.05):r[0].text=f'{i:02d}'
            elif t.isdigit() and abs(sh.left-Inches(12.15))<Inches(.05):r[0].text=str(i)

def main():
    a=argparse.ArgumentParser();a.add_argument('--input',default=str(root.parent/'Revisiting_Lifetime_Aware_VM_Allocation_Presentation_Updated.pptx'))
    a.add_argument('--output',default=str(root.parent/'Revisiting_Lifetime_Aware_VM_Allocation_Presentation_v3.pptx'));args=a.parse_args()
    prs=Presentation(args.input);old=list(prs.slides)
    pics=[next(sh for sh in old[i].shapes if sh.shape_type==13).image.blob for i in (0,1)]  # title graphic, placement figure
    lst=prs.slides._sldIdLst
    for sid in list(lst):prs.part.drop_rel(sid.rId);lst.remove(sid)
    builders=[lambda p:s_title(p,pics[0]),s_overview,lambda p:s_placement(p,pics[1]),s_lava,s_scope,s_methods,s_settings,s_sweep,s_heldout,s_routing,
              s_negative,s_diagnosis,s_method,s_protocol,s_results,s_table,s_decomp,s_robust,s_findings,s_close,s_backup_tau,s_backup_all]
    for b in builders:b(prs)
    prs.part.rename_slide_parts([sid.rId for sid in lst])
    notes=speaker_notes();assert len(notes)==len(prs.slides),(len(notes),len(prs.slides))
    for sl,note in zip(prs.slides,notes):sl.notes_slide.notes_text_frame.text=note
    renumber(prs);prs.save(args.output);print(args.output,len(prs.slides),'slides')
if __name__=='__main__':main()
